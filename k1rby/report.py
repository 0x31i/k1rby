"""k1rby xlsx report — turns the collected sections into a clean, multi-sheet workbook
(ADRecon-style), with a Summary tab and light highlighting of the risky rows."""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="2F3B52")
_RISK_FILL = PatternFill("solid", fgColor="FBE4D5")   # soft amber for risky rows
_TITLE = Font(bold=True, size=14)


def _autosize(ws, rows: list[dict]) -> None:
    if not rows:
        return
    cols = list(rows[0].keys())
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 70)


def _sheet(wb: Workbook, name: str, rows: list[dict], risk_when=None,
           empty_msg: str | None = None) -> int:
    ws = wb.create_sheet(name[:31])
    if not rows:
        cell = ws["A1"]
        cell.value = empty_msg or "(no results)"
        cell.font = Font(italic=True, color="5A6B86")
        return 0
    cols = [c for c in rows[0].keys() if not c.startswith("_")]
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i, c)
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
    for ri, r in enumerate(rows, 2):
        risky = bool(risk_when(r)) if risk_when else False
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(ri, ci, _fmt(r.get(c)))
            if risky:
                cell.fill = _RISK_FILL
    ws.freeze_panes = "A2"
    _autosize(ws, [{c: r.get(c) for c in cols} for r in rows])
    return len(rows)


def _fmt(v: Any) -> Any:
    if isinstance(v, bool):
        return "YES" if v else "no"
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    return "" if v is None else v


def _findings_sheet(wb: Workbook, findings: list) -> None:
    ws = wb.create_sheet("Findings", 0)  # Summary (inserted at 0 later) ends up before it
    cols = ["Severity", "ID", "Title", "Category", "MITRE", "Affected", "Remediation"]
    sev_fill = {"critical": "F2C2C4", "high": "FBD5BF", "medium": "FBE9C5",
                "low": "E3E3E3", "info": "D6E4FB"}
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i, c)
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
    for ri, f in enumerate(findings, 2):
        vals = [f.severity.upper(), f.id, f.title, f.category, f.mitre, f.count, f.remediation]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(ri, ci, v)
            cell.fill = PatternFill("solid", fgColor=sev_fill.get(f.severity, "FFFFFF"))
            cell.alignment = Alignment(wrap_text=(ci == 7), vertical="top")
    ws.freeze_panes = "A2"
    for i, w in enumerate([11, 9, 44, 15, 20, 10, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(path: str, domain: str, sections: dict[str, list[dict]],
          external: dict[str, str] | None = None,
          findings: list | None = None, score: int | None = None) -> dict[str, int]:
    """sections: {sheet_name: rows}. Returns per-sheet counts (also used for the Summary)."""
    wb = Workbook()
    wb.remove(wb.active)

    # risk highlighters per sheet
    risk = {
        "Users": lambda r: r.get("asrepRoastable") or r.get("pwdNotRequired")
        or r.get("unconstrainedDeleg") or (r.get("hasSPN") and str(r.get("adminCount")) == "1"),
        "Computers": lambda r: r.get("unconstrainedDeleg") or r.get("rbcd"),
        "Kerberoastable": lambda r: str(r.get("adminCount")) == "1",
        "AS-REP Roastable": lambda r: True,
        "Delegation": lambda r: "unconstrained" in str(r.get("delegation", "")),
        "Privileged Users": lambda r: r.get("pwdNeverExpires"),
        "LAPS": lambda r: bool(r.get("lapsPassword")),
    }

    # Explain *why* a security-sensitive tab is empty — "clean" vs "not scanned" ambiguity.
    n_users = len(sections.get("Users", []))
    n_comp = len(sections.get("Computers", []))
    empty_msgs = {
        "Kerberoastable":
            f"No kerberoastable accounts — checked {n_users} users; none have a service SPN "
            f"(krbtgt is excluded by design). A clean result.",
        "AS-REP Roastable":
            f"No AS-REP-roastable accounts — checked {n_users} users; none have Kerberos "
            f"pre-authentication disabled (DONT_REQ_PREAUTH). A clean result.",
        "LAPS":
            f"No readable LAPS passwords — checked {n_comp} computers; LAPS is either not "
            f"deployed or not readable by this account. (Absence here is not proof LAPS is absent.)",
        "Trusts":
            "No domain or forest trusts are configured.",
        "Delegation":
            f"No Kerberos delegation configured — checked {n_users} users and {n_comp} "
            f"computers; none have unconstrained, constrained, or resource-based delegation.",
        "Password Policies (FGPP)":
            "No fine-grained password policies defined — only the default domain policy applies "
            "(see the Domain and Password Policy (Compliance) tabs).",
        "Privileged Users":
            f"No accounts flagged adminCount=1 — checked {n_users} users.",
    }

    counts: dict[str, int] = {}
    for name, rows in sections.items():
        counts[name] = _sheet(wb, name, rows, risk.get(name), empty_msg=empty_msgs.get(name))

    # ---- Findings tab (after Summary) ----
    if findings:
        _findings_sheet(wb, findings)

    # ---- Summary tab (first) ----
    summ = wb.create_sheet("Summary", 0)
    summ["A1"] = f"Active Directory Assessment — {domain}"
    summ["A1"].font = _TITLE
    if score is not None:
        summ["A2"] = f"Posture score: {score}/100   ·   {len(findings or [])} findings"
        summ["A2"].font = Font(bold=True, size=12,
                               color=("C00000" if score < 60 else "B36B00" if score < 80 else "2E7D32"))
    summ["A3"] = "Section"
    summ["B3"] = "Count"
    summ["A3"].font = summ["B3"].font = _HEAD
    summ["A3"].fill = summ["B3"].fill = _HEAD_FILL
    r = 4
    for name, n in counts.items():
        summ.cell(r, 1, name)
        summ.cell(r, 2, n)
        r += 1
    if external:
        r += 1
        summ.cell(r, 1, "External collectors").font = Font(bold=True)
        r += 1
        for tool, status in external.items():
            summ.cell(r, 1, tool)
            summ.cell(r, 2, status)
            r += 1
    summ.column_dimensions["A"].width = 34
    summ.column_dimensions["B"].width = 50
    summ["A1"].alignment = Alignment(vertical="center")

    wb.save(path)
    return counts
